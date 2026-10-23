#!/usr/bin/env python3
"""JSON Deepgram -> оформленный docx (титул, таймкоды [ч:мм:сс], номера страниц).

python make_docx.py <расшифровка.json> <выход.docx> "<Название>" ["<Подзаголовок>"]
Нужно: pip install python-docx
"""
import json, sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ACC, GR = RGBColor(0x8A, 0x3B, 0x5C), RGBColor(0x88, 0x88, 0x88)


def ts(t):
    t = int(t)
    return f"{t // 3600}:{t % 3600 // 60:02d}:{t % 60:02d}"


def build(src, out, title, sub=""):
    j = json.load(open(src, encoding="utf8"))
    a = j["results"]["channels"][0]["alternatives"][0]
    d = Document()
    s = d.sections[0]
    s.left_margin = s.right_margin = Cm(2.5)
    s.top_margin = s.bottom_margin = Cm(2.2)
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Georgia", Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Georgia")
    st.paragraph_format.space_after = Pt(8)
    st.paragraph_format.line_spacing = 1.3
    for n, sz in (("Title", 30), ("Heading 1", 20)):
        h = d.styles[n]
        h.font.name, h.font.size, h.font.bold = "Georgia", Pt(sz), False
        h.font.color.rgb = ACC
        for k in ("w:ascii", "w:hAnsi", "w:eastAsia"):
            h.element.rPr.rFonts.set(qn(k), "Georgia")
    d.add_paragraph(style="Title").add_run("Расшифровка: " + title)
    if sub:
        r = d.add_paragraph().add_run(sub)
        r.italic, r.font.size, r.font.color.rgb = True, Pt(13), GR
    r = d.add_paragraph().add_run(
        "Автоматическая расшифровка (Deepgram, русский). Таймкоды стоят в начале абзацев. "
        "Имена, термины и числа стоит проверить на слух: распознаватель может ошибаться.")
    r.font.size, r.font.color.rgb = Pt(10), GR
    d.add_page_break()
    d.add_heading(title, level=1)
    r = d.add_paragraph().add_run("Длительность " + ts(j["metadata"]["duration"]))
    r.font.size, r.font.color.rgb = Pt(10), GR
    for para in a["paragraphs"]["paragraphs"]:
        p = d.add_paragraph()
        r = p.add_run(f"[{ts(para['start'])}]  ")
        r.font.size, r.font.color.rgb, r.bold = Pt(9), ACC, True
        p.add_run(" ".join(x["text"] for x in para["sentences"]))
    fp = s.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tag)
        else:
            e = OxmlElement("w:instrText"); e.text = txt
        run._r.append(e)
    d.save(out)


if __name__ == "__main__":
    if len(sys.argv) < 4:
        sys.exit(__doc__)
    build(*sys.argv[1:5])
